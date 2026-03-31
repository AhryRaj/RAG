from docx import Document
import os
import glob

def debug_docx_tables():
    # Find the file - handle potential naming variations
    files = glob.glob("uploaded_documents/*BRD_Product*vv3*.docx")
    if not files:
        # Try finding *any* BRD_Product file
        files = glob.glob("uploaded_documents/*BRD_Product*.docx")
    
    if not files:
        print("[FAIL] No BRD_Product DOCX file found in uploaded_documents/")
        return
        
    filepath = files[0]
    print(f"[OK] Analyzing file: {filepath}")
    
    try:
        doc = Document(filepath)
        
        print("\n--- Checking Paragraphs ---")
        param_match = False
        for p in doc.paragraphs:
            if "non-functional" in p.text.lower():
                print(f"[MATCH] Found in Paragraph: {p.text[:50]}...")
                print(f"        Style: '{p.style.name}'")
                param_match = True
        if not param_match:
            print("[FAIL] Not found in paragraphs.")

        print("\n--- Checking Tables ---")
        table_match = False
        for i, table in enumerate(doc.tables):
            for row in table.rows:
                for cell in row.cells:
                    if "non-functional" in cell.text.lower():
                        clean_text = cell.text.replace('\n', ' ')
                        print(f"[MATCH] Found in Table {i+1}: {clean_text[:60]}...")
                        table_match = True
        if not table_match:
            print("[FAIL] Not found in tables.")
            
    except Exception as e:
        print(f"[ERROR] {e}")

if __name__ == "__main__":
    debug_docx_tables()
